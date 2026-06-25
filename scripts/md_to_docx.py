#!/usr/bin/env python3
"""
将Markdown文件转换为Word文档
支持标题、表格、列表、粗体、斜体等格式
"""
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def set_chinese_font(run, font_name='宋体', size=12, bold=False):
    """设置中文字体"""
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.font.bold = bold

def parse_markdown(md_text):
    """解析Markdown文本为结构化数据"""
    lines = md_text.split('\n')
    result = []
    i = 0
    
    while i < len(lines):
        line = lines[i]
        
        # 跳过空行
        if not line.strip():
            i += 1
            continue
        
        # 标题
        if line.startswith('# '):
            result.append(('h1', line[2:].strip()))
        elif line.startswith('## '):
            result.append(('h2', line[3:].strip()))
        elif line.startswith('### '):
            result.append(('h3', line[4:].strip()))
        elif line.startswith('#### '):
            result.append(('h4', line[5:].strip()))
        
        # 表格
        elif line.startswith('|') and i + 1 < len(lines) and lines[i + 1].startswith('|'):
            table_lines = [line]
            i += 1
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            result.append(('table', parse_table(table_lines)))
            continue
        
        # 代码块
        elif line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code_lines.append(lines[i])
                i += 1
            result.append(('code', '\n'.join(code_lines)))
        
        # 列表项
        elif re.match(r'^[\s]*[\-\*]\s', line):
            result.append(('ul', re.sub(r'^[\s]*[\-\*]\s', '', line)))
        elif re.match(r'^[\s]*\d+\.\s', line):
            result.append(('ol', re.sub(r'^[\s]*\d+\.\s', '', line)))
        
        # 普通段落
        else:
            result.append(('p', line))
        
        i += 1
    
    return result

def parse_table(table_lines):
    """解析表格"""
    rows = []
    for i, line in enumerate(table_lines):
        # 跳过分隔行（如 |:---:|:---:|
        if re.match(r'^[\s]*\|[-:\s|]+\|[\s]*$', line):
            continue
        
        cells = [cell.strip() for cell in line.split('|')[1:-1]]
        if cells:
            rows.append(cells)
    return rows

def add_formatted_text(paragraph, text):
    """添加带格式的文本（粗体、斜体）"""
    # 处理粗体 **text**
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_chinese_font(run, '宋体', 12, True)
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            set_chinese_font(run, '宋体', 12)
        else:
            run = paragraph.add_run(part)
            set_chinese_font(run, '宋体', 12)

def md_to_docx(md_file, docx_file):
    """转换Markdown为Word文档"""
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    style.font.name = '宋体'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.font.size = Pt(12)
    
    # 读取Markdown
    with open(md_file, 'r', encoding='utf-8') as f:
        md_text = f.read()
    
    # 解析Markdown
    elements = parse_markdown(md_text)
    
    # 添加到文档
    for elem_type, content in elements:
        if elem_type == 'h1':
            p = doc.add_heading(content, level=1)
            for run in p.runs:
                set_chinese_font(run, '黑体', 16, True)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        elif elem_type == 'h2':
            p = doc.add_heading(content, level=2)
            for run in p.runs:
                set_chinese_font(run, '黑体', 14, True)
            
        elif elem_type == 'h3':
            p = doc.add_heading(content, level=3)
            for run in p.runs:
                set_chinese_font(run, '黑体', 12, True)
            
        elif elem_type == 'h4':
            p = doc.add_heading(content, level=4)
            for run in p.runs:
                set_chinese_font(run, '黑体', 12)
            
        elif elem_type == 'p':
            p = doc.add_paragraph()
            add_formatted_text(p, content)
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.line_spacing = 1.5
            
        elif elem_type == 'ul':
            p = doc.add_paragraph(content, style='List Bullet')
            for run in p.runs:
                set_chinese_font(run, '宋体', 12)
            
        elif elem_type == 'ol':
            p = doc.add_paragraph(content, style='List Number')
            for run in p.runs:
                set_chinese_font(run, '宋体', 12)
            
        elif elem_type == 'code':
            p = doc.add_paragraph()
            run = p.add_run(content)
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
            p.paragraph_format.left_indent = Inches(0.5)
            
        elif elem_type == 'table':
            if content:
                table = doc.add_table(rows=len(content), cols=len(content[0]))
                table.style = 'Table Grid'
                for i, row in enumerate(content):
                    for j, cell_text in enumerate(row):
                        cell = table.rows[i].cells[j]
                        cell.text = cell_text
                        # 设置字体
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                set_chinese_font(run, '宋体', 10)
                doc.add_paragraph()  # 表格后空行
    
    # 保存文档
    doc.save(docx_file)
    print(f"已保存: {docx_file}")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("用法: python md_to_docx.py <markdown文件> [输出docx文件]")
        sys.exit(1)
    
    md_file = sys.argv[1]
    if len(sys.argv) >= 3:
        docx_file = sys.argv[2]
    else:
        docx_file = md_file.replace('.md', '.docx')
    
    md_to_docx(md_file, docx_file)
